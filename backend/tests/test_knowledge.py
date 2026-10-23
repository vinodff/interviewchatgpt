import asyncio
import base64
import io
import json
import stat
import zipfile
from datetime import UTC, datetime, timedelta

import httpx
import pytest
import pytest_asyncio
from PIL import Image
from pypdf import PdfWriter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.models import Base, Chunk, Project, Source, User
from app.services import extraction, knowledge, worker
from app.services.extraction import ExtractionError, Segment, extract_upload, parse_github_url


def archive_bytes(entries: dict[str, bytes], compression=zipfile.ZIP_STORED) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=compression) as archive:
        for name, content in entries.items():
            entry = zipfile.ZipInfo("placeholder")
            # ZipInfo normalizes Windows separators; set the raw malicious filename afterward.
            entry.filename = name
            entry.compress_type = compression
            archive.writestr(entry, content)
    return output.getvalue()


class FakeAI:
    def __init__(self, fail=False):
        self.fail = fail
        self.files = []

    async def embed(self, texts, task_type):
        if self.fail:
            raise RuntimeError("secret=must-not-be-in-source-status")
        return [[1.0] + [0.0] * 767 for _ in texts]

    async def structured(self, prompt, schema, fast=False):
        evidence = next(json.loads(line) for line in prompt.splitlines() if line.startswith('{"locator"'))
        return {
            "facts": [
                {
                    "category": "technology",
                    "label": "Python",
                    "quote": evidence["text"],
                    "locator": evidence["locator"],
                    "classification": "reported_experience",
                }
            ]
        }

    async def analyze_file(self, data, mime, prompt):
        self.files.append((data, mime))
        return "I implemented a Python customer API."

    async def close(self):
        pass


@pytest_asyncio.fixture
async def database(monkeypatch):
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    monkeypatch.setattr(knowledge, "SessionLocal", factory)
    monkeypatch.setattr(worker, "SessionLocal", factory)
    monkeypatch.setattr(knowledge, "GeminiService", FakeAI)
    async with factory() as db:
        db.add_all(
            [
                User(id="alice", email="alice@example.com", name="Alice", password_hash="test"),
                User(id="bob", email="bob@example.com", name="Bob", password_hash="test"),
            ]
        )
        await db.commit()
    yield factory
    await engine.dispose()


@pytest.mark.parametrize(
    "url",
    [
        "http://github.com/org/repo",
        "https://github.com.evil.test/org/repo",
        "https://github.com@evil.test/org/repo",
        "https://github.com:443/org/repo",
        "https://github.com/org/repo/tree/main",
        "https://github.com/org/repo?url=evil",
        "https://api.github.com/org/repo",
        "https://github.com/org/..",
        "https://github.com/org/repo#main",
        "https://github.com/org%2fescape/repo",
        "https://github.com/org//repo",
        "https://github.com//org/repo",
        "https://github.com/org/repo//",
    ],
)
def test_rejects_noncanonical_github_urls(url):
    with pytest.raises(ExtractionError):
        parse_github_url(url)


def test_parses_repository_only():
    assert parse_github_url("https://github.com/example/a-project.git/") == ("example", "a-project")


@pytest.mark.parametrize(
    "path",
    ["../escape.py", "/root.py", "C:/escape.py", "a\\escape.py", ".env", "keys/id_rsa", "certs/private.pem"],
)
async def test_archive_rejects_unsafe_or_secret_paths(path):
    with pytest.raises(ExtractionError):
        await extract_upload(
            archive_bytes({path: b"print('hello')"}), "project.zip", "application/zip", FakeAI()
        )


async def test_archive_rejects_symlink_and_high_compression_ratio():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        link = zipfile.ZipInfo("link.py")
        link.create_system = 3
        link.external_attr = (stat.S_IFLNK | 0o777) << 16
        archive.writestr(link, "../outside")
    with pytest.raises(ExtractionError, match="symbolic link"):
        await extract_upload(output.getvalue(), "project.zip", "application/zip", FakeAI())
    compressed = archive_bytes({"large.txt": b"a" * 200_000}, zipfile.ZIP_DEFLATED)
    with pytest.raises(ExtractionError, match="compression-ratio"):
        await extract_upload(compressed, "project.zip", "application/zip", FakeAI())


async def test_archive_rejects_entry_limit_and_case_collision(monkeypatch):
    monkeypatch.setattr(extraction, "MAX_FILES", 2)
    with pytest.raises(ExtractionError, match="entry limit"):
        await extract_upload(
            archive_bytes({"a.py": b"x", "b.py": b"x", "c.py": b"x"}),
            "project.zip",
            "application/zip",
            FakeAI(),
        )
    with pytest.raises(ExtractionError, match="duplicate"):
        await extract_upload(
            archive_bytes({"App.py": b"x", "app.py": b"y"}), "project.zip", "application/zip", FakeAI()
        )


async def test_archive_covers_layers_and_never_executes_code():
    result = await extract_upload(
        archive_bytes(
            {
                "README.md": b"Python FastAPI backend and a React frontend.",
                "backend/api.py": b"raise RuntimeError('this must never execute')",
                "frontend/page.tsx": b"export default function Page() { return <h1>Hello</h1> }",
                "database/schema.sql": b"CREATE TABLE customers(id INTEGER PRIMARY KEY);",
                "binary.py": b"\x00\x01binary",
                "node_modules/dependency.js": b"ignored dependency",
            }
        ),
        "project.zip",
        "application/zip",
        FakeAI(),
    )
    locators = [segment.locator for segment in result.segments]
    assert any("README" in locator for locator in locators)
    assert any("backend/api" in locator for locator in locators)
    assert any("frontend/page" in locator for locator in locators)
    assert any("database/schema" in locator for locator in locators)
    assert not any("node_modules" in locator or "binary.py" in locator for locator in locators)
    assert result.metadata["executed_code"] is False
    assert any("binary" in warning for warning in result.warnings)


async def test_archive_reports_partial_coverage(monkeypatch):
    monkeypatch.setattr(extraction, "MAX_SELECTED_FILES", 2)
    result = await extract_upload(
        archive_bytes({f"file{number}.py": b"import math" for number in range(4)}),
        "project.zip",
        "application/zip",
        FakeAI(),
    )
    assert any("partial" in warning for warning in result.warnings)
    assert result.metadata["files_selected"] == 2


def test_sampling_keeps_all_layers_with_many_migrations():
    paths = [f"database/migration_{number}.sql" for number in range(150)]
    paths += ["README.md", "package.json", "backend/api.py", "frontend/page.tsx", "docs/architecture.md"]
    selected = extraction.select_repository_paths(paths)
    assert len(selected) == 80
    assert all(
        path in selected
        for path in [
            "README.md",
            "package.json",
            "backend/api.py",
            "frontend/page.tsx",
            "docs/architecture.md",
        ]
    )
    assert extraction.select_repository_paths(list(reversed(paths))) == selected


def test_analysis_sample_round_robins_file_chunks(monkeypatch):
    monkeypatch.setattr(knowledge, "MAX_ANALYSIS_CHARS", 400)
    segments = [Segment("a" * 100, f"README.md:L{number}") for number in range(10)]
    segments += [Segment("b" * 100, "frontend/page.tsx:L1")]
    sample = [json.loads(segment) for segment in knowledge._analysis_sample(segments)]
    assert any(segment["locator"] == "frontend/page.tsx:L1" for segment in sample)


async def test_text_locators_and_secret_redaction():
    result = await extract_upload(
        b"I developed Python APIs.\n-----BEGIN PRIVATE KEY-----\nSECRET-BODY\n-----END PRIVATE KEY-----\n",
        "resume.txt",
        "text/plain",
        FakeAI(),
    )
    assert result.segments[0].locator == "resume.txt:L1-L2"
    assert "[REDACTED SECRET]" in result.segments[0].text
    assert "SECRET-BODY" not in result.segments[0].text
    with pytest.raises(ExtractionError, match="Binary"):
        await extract_upload(b"\x00binary", "file.py", "text/plain", FakeAI())


async def test_docx_paragraph_and_table_locators():
    from docx import Document

    document = Document()
    document.add_paragraph("I built a Python service.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "React frontend"
    output = io.BytesIO()
    document.save(output)
    result = await extract_upload(
        output.getvalue(),
        "resume.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        FakeAI(),
    )
    assert any(segment.locator.startswith("paragraph 1") for segment in result.segments)
    assert any(segment.locator.startswith("table 1, row 1") for segment in result.segments)


async def test_images_and_scanned_pdf_use_provider_with_page_locators():
    output = io.BytesIO()
    Image.new("RGB", (10, 10), "white").save(output, format="PNG")
    ai = FakeAI()
    image = await extract_upload(output.getvalue(), "resume.png", "image/png", ai)
    assert image.segments[0].locator.startswith("image 1")
    assert ai.files[0][1] == "image/png"
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    pdf.write(output)
    scanned = await extract_upload(output.getvalue(), "resume.pdf", "application/pdf", ai)
    assert scanned.segments[0].locator.startswith("page 1")
    assert ai.files[1][1] == "application/pdf"
    assert scanned.warnings


def test_grounding_rejects_invented_quotes_and_forces_reported_attribution():
    segment = Segment("I developed Python APIs for customer records.", "paragraph 1")
    raw = {
        "facts": [
            {
                "category": "experience",
                "label": "Python APIs",
                "quote": "I developed Python APIs",
                "locator": "paragraph 1",
                "classification": "repository_observation",
            },
            {
                "category": "achievement",
                "label": "Million-dollar savings",
                "quote": "Saved one million dollars",
                "locator": "paragraph 1",
                "classification": "reported_experience",
            },
            {
                "category": "experience",
                "label": "Missing locator",
                "quote": "I developed Python APIs",
                "locator": "page 999",
                "classification": "reported_experience",
            },
        ]
    }
    facts = knowledge._ground_facts(raw, [segment], repository=False)
    assert len(facts) == 1
    assert facts[0]["classification"] == "reported_experience"
    assert facts[0]["verified_contribution"] is False
    assert (
        knowledge._ground_facts(raw, [segment], repository=True)[0]["classification"]
        == "repository_observation"
    )


async def test_github_uses_fixed_host_revision_and_unverified_identity(monkeypatch):
    sha = "a" * 40
    requests = []
    transport = httpx.MockTransport(lambda request: response(request))

    def response(request):
        requests.append(request)
        assert request.url.host == "api.github.com"
        path = request.url.path
        if path.endswith("/commits"):
            data = [
                {
                    "sha": sha,
                    "author": {"login": "alice"},
                    "commit": {"message": "Add Python API"},
                    "url": "http://169.254.169.254/secrets",
                }
            ]
        elif "/git/trees/" in path:
            data = {
                "truncated": True,
                "tree": [
                    {
                        "path": "README.md",
                        "type": "blob",
                        "mode": "100644",
                        "size": 20,
                        "sha": sha,
                        "url": "https://evil.test/blob",
                    },
                    {"path": ".env", "type": "blob", "mode": "100644", "size": 10, "sha": sha},
                ],
            }
        elif "/git/blobs/" in path:
            data = {"encoding": "base64", "content": base64.b64encode(b"A Python API repository.").decode()}
        else:
            data = {"default_branch": "main", "url": "https://evil.test/redirect"}
        return httpx.Response(200, json=data)

    actual_client = httpx.AsyncClient
    monkeypatch.setattr(
        extraction.httpx, "AsyncClient", lambda **kwargs: actual_client(transport=transport, **kwargs)
    )
    result = await extraction.extract_github("https://github.com/alice/project", "alice")
    assert result.metadata["identity_verified"] is False
    assert result.metadata["revision"] == sha
    assert result.metadata["tree_truncated"] is True
    commit = next(segment for segment in result.segments if segment.locator.startswith("commit"))
    assert commit.metadata["declared_username_match"] is True
    assert commit.metadata["verified_contribution"] is False
    assert not any(".env" in str(request.url) for request in requests)


async def test_successful_ingestion_persists_grounded_profile(database, monkeypatch):
    async def uploaded(_path):
        return b"I built a Python customer API."

    monkeypatch.setattr(knowledge, "_read_upload", uploaded)
    async with database() as db:
        db.add(Project(id="alice-project", user_id="alice", name="Customer API", description="My project"))
        await db.flush()
        db.add(
            Source(
                id="alice-source",
                user_id="alice",
                project_id="alice-project",
                name="resume.txt",
                storage_path="private.enc",
                status="queued",
            )
        )
        await db.commit()
    assert await worker.process_one() is True
    async with database() as db:
        source = await db.get(Source, "alice-source")
        user = await db.get(User, "alice")
        project = await db.get(Project, "alice-project")
        chunks = list((await db.scalars(select(Chunk))).all())
        assert source.status == "ready" and source.chunk_count == 1
        assert len(chunks[0].embedding) == 768
        assert chunks[0].verified_contribution is False
        assert user.profile["reported_experience"][0]["locator"] == chunks[0].locator
        assert project.analysis["technologies"] == ["Python"]


async def test_embedding_failure_does_not_mark_ready_and_backoff_is_durable(database, monkeypatch):
    async def uploaded(_path):
        return b"I built a Python customer API."

    monkeypatch.setattr(knowledge, "_read_upload", uploaded)
    monkeypatch.setattr(knowledge, "GeminiService", lambda: FakeAI(fail=True))
    async with database() as db:
        db.add(
            Source(
                id="failed-source",
                user_id="alice",
                name="resume.txt",
                storage_path="private.enc",
                status="queued",
            )
        )
        await db.commit()
    await worker.process_one()
    async with database() as db:
        source = await db.get(Source, "failed-source")
        assert source.status == "queued"
        assert source.attempts == 1
        assert source.analysis["_worker"]["next_attempt_at"]
        assert "secret=" not in source.error
        assert not list((await db.scalars(select(Chunk))).all())
    assert await worker.claim_source() is None


async def test_upload_integrity_failure_never_publishes_evidence(database, monkeypatch):
    async def uploaded(_path):
        return b"Changed content pretending to be the original resume."

    monkeypatch.setattr(knowledge, "_read_upload", uploaded)
    async with database() as db:
        db.add(
            Source(
                id="tampered",
                user_id="alice",
                name="resume.txt",
                storage_path="private.enc",
                status="queued",
                sha256="0" * 64,
            )
        )
        await db.commit()
    await worker.process_one()
    async with database() as db:
        source = await db.get(Source, "tampered")
        assert source.status == "queued"
        assert "integrity" in source.error
        assert not list((await db.scalars(select(Chunk))).all())


async def test_third_attempt_failure_is_terminal(database):
    async with database() as db:
        db.add(
            Source(
                id="last-attempt",
                user_id="alice",
                name="resume.txt",
                status="processing",
                attempts=3,
                processing_started_at=datetime.now(UTC),
                analysis={"_worker": {"lease_token": "last"}},
            )
        )
        await db.commit()
    await worker.mark_failure("last-attempt", "last", RuntimeError("secret=upstream"))
    async with database() as db:
        source = await db.get(Source, "last-attempt")
        assert source.status == "failed"
        assert "secret=" not in source.error
    assert await worker.claim_source() is None


async def test_retry_limit_and_stale_processing_recovery(database):
    async with database() as db:
        db.add_all(
            [
                Source(
                    id="stale",
                    user_id="alice",
                    name="old.txt",
                    status="processing",
                    attempts=1,
                    processing_started_at=datetime.now(UTC) - timedelta(minutes=20),
                    analysis={"_worker": {"lease_token": "old"}},
                ),
                Source(
                    id="exhausted",
                    user_id="alice",
                    name="failed.txt",
                    status="processing",
                    attempts=3,
                    processing_started_at=datetime.now(UTC) - timedelta(minutes=20),
                ),
            ]
        )
        await db.commit()
    claim = await worker.claim_source()
    assert claim[0] == "stale" and claim[1] != "old"
    async with database() as db:
        stale = await db.get(Source, "stale")
        assert stale.attempts == 2 and stale.status == "processing"
        assert (await db.get(Source, "exhausted")).status == "failed"
    await worker.mark_failure("stale", "old", RuntimeError("old worker"))
    async with database() as db:
        assert (await db.get(Source, "stale")).status == "processing"


async def test_lost_lease_discards_extracted_chunks(database, monkeypatch):
    async def uploaded(_path):
        async with database() as db:
            source = await db.get(Source, "lease-source")
            source.analysis = {"_worker": {"lease_token": "replacement"}}
            await db.commit()
        return b"I built a Python customer API."

    monkeypatch.setattr(knowledge, "_read_upload", uploaded)
    async with database() as db:
        db.add(
            Source(
                id="lease-source",
                user_id="alice",
                name="resume.txt",
                storage_path="private.enc",
                status="processing",
                processing_started_at=datetime.now(UTC),
                analysis={"_worker": {"lease_token": "original"}},
            )
        )
        await db.commit()
    with pytest.raises(knowledge.LeaseLost):
        await knowledge.ingest_source("lease-source")
    async with database() as db:
        assert not list((await db.scalars(select(Chunk))).all())
        assert (await db.get(Source, "lease-source")).status == "processing"


async def test_retrieval_scopes_both_source_and_chunk_owner_and_project(database, monkeypatch):
    async with database() as db:
        db.add_all([Project(id="pa", user_id="alice", name="A"), Project(id="pb", user_id="bob", name="B")])
        await db.flush()
        db.add_all(
            [
                Source(id="sa", user_id="alice", project_id="pa", name="alice.py", status="ready"),
                Source(id="sb", user_id="bob", project_id="pb", name="bob.py", status="ready"),
                Source(id="pending", user_id="alice", name="pending.txt", status="processing"),
            ]
        )
        await db.flush()
        for identifier, owner, source, project in [
            ("ca", "alice", "sa", "pa"),
            ("cb", "bob", "sb", "pb"),
            ("corrupt", "alice", "sb", "pb"),
            ("unready", "alice", "pending", None),
        ]:
            db.add(
                Chunk(
                    id=identifier,
                    user_id=owner,
                    source_id=source,
                    project_id=project,
                    text="Python API customer records",
                    locator="L1",
                    embedding=[1.0] + [0.0] * 767,
                )
            )
        await db.commit()
    found = await knowledge.retrieve("alice", "Python API", project_id="pa")
    assert [evidence["id"] for evidence in found] == ["ca"]
    assert await knowledge.retrieve("alice", "Python API", project_id="pb") == []
    monkeypatch.setattr(knowledge, "GeminiService", lambda: FakeAI(fail=True))
    lexical = await knowledge.retrieve("alice", "Python API")
    assert [evidence["id"] for evidence in lexical] == ["ca"]
    assert lexical[0]["retrieval_mode"] == "lexical_fallback"


@pytest.mark.parametrize("vector", [[0.0] * 768, [1.0] * 767, [float("nan")] * 768, [float("inf")] * 768])
def test_invalid_embeddings_are_rejected(vector):
    with pytest.raises(knowledge.IngestionError):
        knowledge._validate_vectors([vector], 1)


async def test_slow_query_embedding_reaches_lexical_fallback_within_live_budget(database, monkeypatch):
    class SlowAI(FakeAI):
        async def embed(self, texts, task_type):
            await asyncio.Event().wait()

    async with database() as db:
        db.add(Source(id="slow-source", user_id="alice", name="notes.txt", status="ready"))
        await db.flush()
        db.add(Chunk(id="slow-chunk", source_id="slow-source", user_id="alice", text="Redis caching reduces database work", locator="L1"))
        await db.commit()
    monkeypatch.setattr(knowledge, "GeminiService", SlowAI)
    async with asyncio.timeout(8):
        evidence = await knowledge.retrieve("alice", "Redis caching")
    assert evidence[0]["id"] == "slow-chunk"
    assert evidence[0]["retrieval_mode"] == "lexical_fallback"
